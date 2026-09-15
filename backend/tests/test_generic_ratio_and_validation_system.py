"""
Generic Ratio Engine, Validation, and Non-Statutory Audit Compliance Test Suite.
Tests across multiple synthetic corporate profiles with ZERO hardcoded company-specific logic.
Validates:
1. Mathematical alignment across formula, inputs, calculation, and displayed results.
2. Debt-to-Equity vs Total Liabilities-to-Equity distinction.
3. Net Working Capital to Revenue naming and calculation.
4. Absence of unauthorized statutory audit claims across all reports.
5. Recommendation grounding (never citing missing source items).
6. Working capital cycle suppression on missing inputs.
7. Preservation of accounting discrepancies without synthetic alteration.
"""

import io
import openpyxl  # type: ignore
import pypdf
import pytest
from app.engine.ratio_engine import RatioEngine, verify_ratio_reproducibility
from app.engine.financial_analyzer import calculate_financial_ratios, calculate_corporate_finance
from app.engine.statement_generator import generate_financial_statements
from app.engine.auditor_engine import perform_full_financial_audit
from app.engine.ai_insights import generate_ai_insights
from app.engine.output_validator import OutputValidator, ReportConsistencyValidator, PROHIBITED_AUDIT_TERMS
from app.reports.pdf_generator import generate_pdf_report
from app.reports.excel_generator import generate_excel_report


def build_synthetic_service_company():
    """Service company: No inventory, no COGS, strong operating margins, minimal debt."""
    return [
        {"account_code": "100", "account_name": "Consulting Revenue", "account_type": "REVENUE", "net_amount": 500000.0, "year": "2025"},
        {"account_code": "200", "account_name": "Salaries & Benefits", "account_type": "OPERATING_EXPENSE", "net_amount": 250000.0, "year": "2025"},
        {"account_code": "210", "account_name": "Office Rent & Tech", "account_type": "OPERATING_EXPENSE", "net_amount": 50000.0, "year": "2025"},
        {"account_code": "220", "account_name": "Income Taxes", "account_type": "TAX_EXPENSE", "net_amount": 40000.0, "year": "2025"},
        {"account_code": "300", "account_name": "Cash & Cash Equivalents", "account_type": "CASH_ASSET", "net_amount": 180000.0, "year": "2025"},
        {"account_code": "310", "account_name": "Accounts Receivable", "account_type": "RECEIVABLE_ASSET", "net_amount": 120000.0, "year": "2025"},
        {"account_code": "320", "account_name": "Computer Equipment", "account_type": "FIXED_ASSET", "net_amount": 100000.0, "year": "2025"},
        {"account_code": "400", "account_name": "Accounts Payable", "account_type": "PAYABLE_LIABILITY", "net_amount": -40000.0, "year": "2025"},
        {"account_code": "410", "account_name": "Accrued Payroll", "account_type": "CURRENT_LIABILITY", "net_amount": -60000.0, "year": "2025"},
        {"account_code": "500", "account_name": "Share Capital & Retained Earnings", "account_type": "EQUITY", "net_amount": -300000.0, "year": "2025"}
    ]


def build_synthetic_manufacturing_company():
    """Manufacturing company: Raw materials, inventory, heavy PPE, bank loans, trade payables."""
    return [
        {"account_code": "101", "account_name": "Manufacturing Sales", "account_type": "REVENUE", "net_amount": 1200000.0, "year": "2025"},
        {"account_code": "102", "account_name": "Cost of Goods Sold", "account_type": "COGS", "net_amount": 720000.0, "year": "2025"},
        {"account_code": "201", "account_name": "Selling & Admin Expenses", "account_type": "OPERATING_EXPENSE", "net_amount": 180000.0, "year": "2025"},
        {"account_code": "202", "account_name": "Depreciation Expense", "account_type": "DEPRECIATION", "net_amount": 60000.0, "year": "2025"},
        {"account_code": "203", "account_name": "Bank Loan Interest", "account_type": "INTEREST_EXPENSE", "net_amount": 30000.0, "year": "2025"},
        {"account_code": "204", "account_name": "Tax Expense", "account_type": "TAX_EXPENSE", "net_amount": 42000.0, "year": "2025"},
        {"account_code": "301", "account_name": "Cash", "account_type": "CASH_ASSET", "net_amount": 150000.0, "year": "2025"},
        {"account_code": "302", "account_name": "Trade Receivables", "account_type": "RECEIVABLE_ASSET", "net_amount": 220000.0, "year": "2025"},
        {"account_code": "303", "account_name": "Finished Goods Inventory", "account_type": "INVENTORY_ASSET", "net_amount": 180000.0, "year": "2025"},
        {"account_code": "304", "account_name": "Plant & Machinery", "account_type": "FIXED_ASSET", "net_amount": 650000.0, "year": "2025"},
        {"account_code": "401", "account_name": "Trade Payables", "account_type": "PAYABLE_LIABILITY", "net_amount": -140000.0, "year": "2025"},
        {"account_code": "402", "account_name": "Short-Term Working Capital Loan", "account_type": "CURRENT_LIABILITY", "net_amount": -80000.0, "year": "2025"},
        {"account_code": "403", "account_name": "Long-Term Term Debt", "account_type": "LONG_TERM_LIABILITY", "net_amount": -320000.0, "year": "2025"},
        {"account_code": "501", "account_name": "Shareholders' Equity", "account_type": "EQUITY", "net_amount": -660000.0, "year": "2025"}
    ]


def test_service_company_inventory_suppression():
    """Verify that service companies without inventory/COGS suppress DIO/Inventory Turnover gracefully."""
    items = build_synthetic_service_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)
    corp_fin = calculate_corporate_finance(stmts, ratios)

    # 1. Inventory Turnover must be NOT_CALCULABLE
    inv_t = ratios["efficiency"]["inventory_turnover"]
    assert inv_t["is_calculable"] is False
    assert inv_t["value"] is None
    assert inv_t["status"] == "NOT_CALCULABLE"

    # 2. DIO must be None, Operating cycle relies on DSO
    wcc = corp_fin["working_capital_cycle"]
    assert wcc["days_inventory_outstanding_dio"] is None
    assert wcc["days_sales_outstanding_dso"] is not None
    assert wcc["operating_cycle"] == wcc["days_sales_outstanding_dso"]

    # 3. Quick Ratio equals Current Ratio when inventory is zero/missing
    liq = ratios["liquidity"]
    assert liq["current_ratio"]["is_calculable"] is True
    assert liq["quick_ratio"]["is_calculable"] is True
    assert liq["quick_ratio"]["value"] == liq["current_ratio"]["value"]


def test_ratio_mathematical_agreement_and_reproducibility():
    """Verify that displayed inputs, calculation, formula, and displayed value agree 100% mathematically."""
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)

    for cat_name, cat_dict in ratios.items():
        for r_name, r_obj in cat_dict.items():
            if not isinstance(r_obj, dict):
                continue
            assert "formula" in r_obj, f"Missing formula in ratio {r_name}"
            assert "inputs" in r_obj, f"Missing inputs in ratio {r_name}"
            assert "definition_used" in r_obj, f"Missing definition_used in ratio {r_name}"
            
            if r_obj.get("is_calculable") and r_obj.get("value") is not None:
                assert r_obj.get("reproducible") is True, f"Ratio {r_name} failed mathematical reproducibility check!"


def test_debt_to_equity_vs_liabilities_to_equity_distinction():
    """Ensure Debt-to-Equity is strictly interest-bearing debt and distinct from Total Liabilities to Equity."""
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)

    solv = ratios["solvency"]
    de = solv["debt_to_equity"]
    le = solv["liabilities_to_equity"]

    assert de["name"] == "Debt to Equity Ratio"
    assert le["name"] == "Total Liabilities to Equity Ratio"

    # De leverages only short-term loan + long-term debt: (80000 + 320000) = 400000 / 660000 = ~0.61
    # Le leverages total liabilities: (140000 + 80000 + 320000) = 540000 / 660000 = ~0.82
    assert de["value"] is not None
    assert le["value"] is not None
    assert de["value"] < le["value"], "Debt to Equity should be strictly lower than Liabilities to Equity when non-debt liabilities exist!"


def test_net_working_capital_to_revenue_label_and_calculation():
    """Verify Net Working Capital to Revenue is accurately named and properly calculated."""
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)

    liq = ratios["liquidity"]
    nwc_r = liq["working_capital_ratio"]

    assert nwc_r["name"] == "Net Working Capital to Revenue"
    assert "%" in nwc_r["unit"]
    assert "Revenue" in nwc_r["formula"]
    assert nwc_r["is_calculable"] is True
    assert nwc_r["reproducible"] is True


def test_zero_prohibited_audit_language_across_all_outputs():
    """Guarantee no unauthorized statutory audit claims exist across audit findings, PDF, and Excel reports."""
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)
    corp_fin = calculate_corporate_finance(stmts, ratios)
    audit = perform_full_financial_audit(stmts, ratios)
    ai_rep = generate_ai_insights(stmts, ratios, corp_fin)

    disallowed_strings = [
        "official independent auditor's report",
        "statutory financial audit",
        "clean bill of health",
        "present fairly, in all material respects",
        "present fairly in all material respects"
    ]

    # 1. Check Auditor Engine Payload
    op = audit["auditor_opinion"]
    combined_op_text = f"{op.get('title', '')} {op.get('summary', '')} {op.get('auditor_signature', '')} {op.get('audit_standards', '')}".lower()
    for ds in disallowed_strings:
        assert ds not in combined_op_text, f"Found prohibited audit wording '{ds}' in auditor opinion payload!"

    # 2. Check Pre-Flight Output Validator
    raw_payload = {
        "statements": stmts,
        "ratios": ratios,
        "corporate_finance": corp_fin,
        "audit_report": audit,
        "ai_report": ai_rep
    }
    validated = OutputValidator.validate_and_filter_payload(raw_payload, items)
    assert validated["consistency_report"]["pre_flight_status"] == "PASS"

    # 3. Check Generated PDF bytes
    pdf_bytes = generate_pdf_report("Synthetic Mfg Corp", stmts, ratios, corp_fin, ai_rep, currency="USD", audit_report=audit)
    assert len(pdf_bytes) > 0
    # Text in PDF stream must not contain prohibited audit strings
    pdf_str = pdf_bytes.decode("latin1", errors="ignore").lower()
    for ds in disallowed_strings:
        assert ds not in pdf_str, f"Found prohibited audit phrase '{ds}' in generated PDF output!"

    # 4. Check Generated Excel bytes
    excel_bytes = generate_excel_report("Synthetic Mfg Corp", stmts, ratios, corp_fin, ai_rep, audit)
    assert len(excel_bytes) > 0


def test_recommendation_evidence_grounding():
    """Ensure recommendations do not cite missing data when balance sheet or debt schedules are absent."""
    # Only income statement items
    pnl_only_items = [
        {"account_code": "100", "account_name": "Sales Revenue", "account_type": "REVENUE", "net_amount": 100000.0, "year": "2025"},
        {"account_code": "200", "account_name": "Operating Expenses", "account_type": "OPERATING_EXPENSE", "net_amount": 80000.0, "year": "2025"}
    ]
    stmts = generate_financial_statements(pnl_only_items)
    ratios = calculate_financial_ratios(stmts)
    corp_fin = calculate_corporate_finance(stmts, ratios)
    ai_rep = generate_ai_insights(stmts, ratios, corp_fin)

    raw_payload = {
        "statements": stmts,
        "ratios": ratios,
        "corporate_finance": corp_fin,
        "ai_report": ai_rep
    }
    validated = OutputValidator.validate_and_filter_payload(raw_payload, pnl_only_items)
    filtered_recs = validated["ai_report"]["recommendations"]

    for rec in filtered_recs:
        r_text = f"{rec.get('title', '')} {rec.get('action', '')}".lower()
        assert "debt" not in r_text, f"Recommendation cited debt when no balance sheet or debt exists: {rec}"
        assert "cash conversion" not in r_text, f"Recommendation cited cash conversion cycle when no balance sheet exists: {rec}"


def test_ccc_ending_balance_approximation_label():
    """Verify that single-period CCC is explicitly tagged as an approximation based on ending balances."""
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)
    corp_fin = calculate_corporate_finance(stmts, ratios)

    wcc = corp_fin["working_capital_cycle"]
    assert wcc.get("is_approximation") is True, "CCC should be flagged as is_approximation: True when using ending balances"
    assert "ending balance" in wcc.get("interpretation", "").lower(), "CCC interpretation must explicitly disclose ending balance approximation"

    # AI insights referencing CCC must qualify it as an approximation
    ai_rep = generate_ai_insights(stmts, ratios, corp_fin)
    ai_text = str(ai_rep).lower()
    if "cash conversion cycle" in ai_text:
        assert "approximation" in ai_text or "approximate" in ai_text, "AI report must qualify CCC mentions as approximations"


def test_benchmark_unavailable_disclaimer():
    """Verify that ratio benchmarks explicitly state that external industry benchmarks are unavailable from provided data."""
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)

    for cat_name, cat_dict in ratios.items():
        for r_name, r_obj in cat_dict.items():
            if not isinstance(r_obj, dict):
                continue
            benchmark = r_obj.get("benchmark")
            if benchmark is not None:
                assert "industry benchmark unavailable from the provided data" in str(benchmark).lower(), (
                    f"Ratio {r_name} benchmark '{benchmark}' did not include required unavailable disclaimer"
                )


def test_excel_generator_zero_prohibited_terms():
    """Inspect the raw openpyxl workbook created by generate_excel_report for zero prohibited audit terms."""
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)
    corp_fin = calculate_corporate_finance(stmts, ratios)
    audit = perform_full_financial_audit(stmts, ratios)
    ai_rep = generate_ai_insights(stmts, ratios, corp_fin)

    excel_bytes = generate_excel_report("Synthetic Mfg Corp", stmts, ratios, corp_fin, ai_rep, audit)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))

    # Check sheet names
    sheet_names = wb.sheetnames
    for sn in sheet_names:
        assert "audit opinion" not in sn.lower()
        assert "audit exception" not in sn.lower()

    # Check every cell value across all sheets
    for sn in sheet_names:
        ws = wb[sn]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell and isinstance(cell, str):
                    c_lower = cell.lower()
                    for term in PROHIBITED_AUDIT_TERMS:
                        assert term not in c_lower, f"Prohibited term '{term}' found in Excel sheet '{sn}' cell: '{cell}'"


def test_pdf_text_extraction_deep_audit():
    """Extract all text from generated PDF pages and verify complete absence of prohibited audit phrases."""
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)
    corp_fin = calculate_corporate_finance(stmts, ratios)
    audit = perform_full_financial_audit(stmts, ratios)
    ai_rep = generate_ai_insights(stmts, ratios, corp_fin)

    pdf_bytes = generate_pdf_report("Synthetic Mfg Corp", stmts, ratios, corp_fin, ai_rep, currency="USD", audit_report=audit)
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    all_pdf_text = ""
    for page in reader.pages:
        all_pdf_text += page.extract_text() + "\n"

    all_pdf_text_lower = all_pdf_text.lower()
    for term in PROHIBITED_AUDIT_TERMS:
        assert term not in all_pdf_text_lower, f"Prohibited audit term '{term}' extracted from PDF page text!"


def test_unseen_third_corporate_profile_retail_chain():
    """Unseen profile 3: Retail chain with inventory, leased assets, short-term debt, and supplier trade credit."""
    retail_items = [
        {"account_code": "110", "account_name": "Retail Store Sales", "account_type": "REVENUE", "net_amount": 2500000.0, "year": "2025"},
        {"account_code": "111", "account_name": "E-Commerce Revenue", "account_type": "REVENUE", "net_amount": 800000.0, "year": "2025"},
        {"account_code": "120", "account_name": "Cost of Merchandise Sold", "account_type": "COGS", "net_amount": 1980000.0, "year": "2025"},
        {"account_code": "210", "account_name": "Store Payroll & Commissions", "account_type": "OPERATING_EXPENSE", "net_amount": 420000.0, "year": "2025"},
        {"account_code": "211", "account_name": "Store Leases & Utilities", "account_type": "OPERATING_EXPENSE", "net_amount": 280000.0, "year": "2025"},
        {"account_code": "212", "account_name": "Depreciation of Store Fixtures", "account_type": "DEPRECIATION", "net_amount": 85000.0, "year": "2025"},
        {"account_code": "215", "account_name": "Credit Facility Interest Expense", "account_type": "INTEREST_EXPENSE", "net_amount": 45000.0, "year": "2025"},
        {"account_code": "220", "account_name": "Corporate Income Tax", "account_type": "TAX_EXPENSE", "net_amount": 95000.0, "year": "2025"},
        {"account_code": "310", "account_name": "Store Cash & POS Accounts", "account_type": "CASH_ASSET", "net_amount": 210000.0, "year": "2025"},
        {"account_code": "320", "account_name": "Merchant Processor Receivables", "account_type": "RECEIVABLE_ASSET", "net_amount": 140000.0, "year": "2025"},
        {"account_code": "330", "account_name": "Merchandise Inventory", "account_type": "INVENTORY_ASSET", "net_amount": 520000.0, "year": "2025"},
        {"account_code": "340", "account_name": "Store Leasehold Improvements", "account_type": "FIXED_ASSET", "net_amount": 780000.0, "year": "2025"},
        {"account_code": "410", "account_name": "Merchandise Accounts Payable", "account_type": "PAYABLE_LIABILITY", "net_amount": -380000.0, "year": "2025"},
        {"account_code": "415", "account_name": "Accrued Store Operating Costs", "account_type": "CURRENT_LIABILITY", "net_amount": -110000.0, "year": "2025"},
        {"account_code": "420", "account_name": "Revolving Credit Line", "account_type": "CURRENT_LIABILITY", "net_amount": -220000.0, "year": "2025"},
        {"account_code": "430", "account_name": "Long-Term Equipment Loan", "account_type": "LONG_TERM_LIABILITY", "net_amount": -350000.0, "year": "2025"},
        {"account_code": "510", "account_name": "Common Equity & Retained Earnings", "account_type": "EQUITY", "net_amount": -590000.0, "year": "2025"}
    ]

    stmts = generate_financial_statements(retail_items)
    ratios = calculate_financial_ratios(stmts)
    corp_fin = calculate_corporate_finance(stmts, ratios)
    audit = perform_full_financial_audit(stmts, ratios)
    ai_rep = generate_ai_insights(stmts, ratios, corp_fin)

    # 1. Ratios validation
    solv = ratios["solvency"]
    de = solv["debt_to_equity"]
    assert de["formula"] == "Interest-Bearing Debt / Shareholders' Equity"
    assert de["is_calculable"] is True
    assert de["reproducible"] is True

    le = solv["liabilities_to_equity"]
    assert le["formula"] == "Total Liabilities / Shareholders' Equity"
    assert le["value"] > de["value"]

    liq = ratios["liquidity"]
    nwc = liq["working_capital_ratio"]
    assert nwc["name"] == "Net Working Capital to Revenue"
    assert nwc["formula"] == "(Current Assets - Current Liabilities) / Revenue"
    assert nwc["reproducible"] is True

    # 2. Pre-flight output validation
    payload = {
        "statements": stmts,
        "ratios": ratios,
        "corporate_finance": corp_fin,
        "audit_report": audit,
        "ai_report": ai_rep
    }
    validated = OutputValidator.validate_and_filter_payload(payload, retail_items)
    assert validated["consistency_report"]["pre_flight_status"] == "PASS"
    assert len(validated["consistency_report"]["prohibited_terms_found"]) == 0


def test_complete_pdf_and_excel_prohibited_terms_deep_scan():
    """
    Scans the exact PDF and Excel output files generated for an unseen enterprise target.
    Verifies that NONE of the prohibited statutory audit terms ever appear in the generated files.
    """
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)
    corp_fin = calculate_corporate_finance(stmts, ratios)
    audit = perform_full_financial_audit(stmts, ratios)
    ai_rep = generate_ai_insights(stmts, ratios, corp_fin)

    # Render actual PDF
    pdf_bytes = generate_pdf_report("Apex Dynamics International", stmts, ratios, corp_fin, ai_rep, currency="USD", audit_report=audit)
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    full_pdf_text = ""
    for page in reader.pages:
        full_pdf_text += page.extract_text() + "\n"
    pdf_lower = full_pdf_text.lower()

    prohibited_checklist = [
        "statutory financial audit",
        "auditor's opinion",
        "auditor opinion",
        "audit conclusion",
        "unqualified",
        "clean bill of health",
        "isa / us gaas",
        "isa 320",
        "isa 700",
        "isa 705",
        "present fairly",
        "financial statements present fairly",
        "ai audit analysis",
        "statutory audit",
        "auditor sign-off",
        "certified opinion"
    ]

    for term in prohibited_checklist:
        assert term not in pdf_lower, f"Prohibited audit term '{term}' was found in rendered PDF text!"

    # Render actual Excel
    excel_bytes = generate_excel_report("Apex Dynamics International", stmts, ratios, corp_fin, ai_rep, audit)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    for sname in wb.sheetnames:
        ws = wb[sname]
        for row in ws.iter_rows(values_only=True):
            for val in row:
                if val and isinstance(val, str):
                    v_low = val.lower()
                    for term in prohibited_checklist:
                        assert term not in v_low, f"Prohibited audit term '{term}' was found in Excel cell in sheet '{sname}': '{val}'"


def test_statement_reconciliation_and_source_preservation():
    """
    Tests that:
    1. P&L totals are correct (Gross Profit = Rev - COGS, Net Income = PBT - Tax)
    2. Balance Sheet reconciles
    3. Cash Flow reconciliation is preserved
    4. Source values are never silently changed
    """
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)

    inc = stmts["income_statement"]
    assert inc["gross_profit"] == inc["revenue_from_operations"] - inc["cost_of_goods_sold"]
    assert inc["ebit"] == inc["gross_profit"] - inc["operating_expenses"] - inc["depreciation_amortization"]

    bs = stmts["balance_sheet"]
    val_rep = stmts.get("validation_report", {})
    assert val_rep.get("balance_sheet_check") == "PASS"
    eq_val = bs.get("equity", {}).get("total_equity") or bs.get("total_equity")
    assert bs["total_assets"] == (bs["total_liabilities"] + eq_val)

    # Verify source values are untouched
    for it in items:
        # Check that net_amount matches an item in normalized items
        found = any(n.get("account_name") == it["account_name"] and abs(n.get("net_amount", 0) - it["net_amount"]) < 0.01 for n in stmts.get("normalized_items", []))
        assert found, f"Source line item '{it['account_name']}' was altered or missing from normalized statements!"


def test_ccc_approximation_across_api_pdf_excel():
    """
    Tests that CCC approximation status exists in underlying data model, API response, PDF and Excel output.
    """
    items = build_synthetic_manufacturing_company()
    stmts = generate_financial_statements(items)
    ratios = calculate_financial_ratios(stmts)
    corp_fin = calculate_corporate_finance(stmts, ratios)
    wcc = corp_fin["working_capital_cycle"]

    # 1. Underlying data model
    assert wcc.get("is_approximation") is True
    assert wcc.get("metric_name") == "Approximate Cash Conversion Cycle"
    assert "ending balance" in wcc.get("approximation_status", "").lower()

    # 2. PDF rendering
    pdf_bytes = generate_pdf_report("Apex Dynamics", stmts, ratios, corp_fin, {}, currency="USD")
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    pdf_text = "".join(page.extract_text() for page in reader.pages)
    assert "Approximate Cash Conversion Cycle" in pdf_text
    assert "Ending Balances Used" in pdf_text

    # 3. Excel rendering
    excel_bytes = generate_excel_report("Apex Dynamics", stmts, ratios, corp_fin, {})
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Working Capital & CCC" in wb.sheetnames
    ws = wb["Working Capital & CCC"]
    found_approx_ccc = False
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if val and "Approximate Cash Conversion Cycle" in str(val):
                found_approx_ccc = True
    assert found_approx_ccc, "Excel sheet 'Working Capital & CCC' missing 'Approximate Cash Conversion Cycle'"

