import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from app.engine.document_parser import parse_workbook, sanitize_json_data
from app.engine.statement_generator import generate_financial_statements
from app.engine.financial_analyzer import calculate_financial_ratios, calculate_corporate_finance
from app.engine.canonical_model import build_canonical_dataset
from app.engine.reconciliation import perform_source_to_result_reconciliation
from app.engine.quality_engine import compute_financial_quality_score
from app.engine.ai_insights import generate_ai_insights
from app.engine.auditor_engine import perform_full_financial_audit

def create_unseen_docx(filepath: str):
    doc = Document()
    
    doc.add_heading("AeroPulse Systems Corp. - Financial Statements", level=1)
    doc.add_paragraph("Fiscal Year: FY 2024–25")
    doc.add_paragraph("Reporting Currency: USD ($) | Unaudited Financial Summary")
    doc.add_paragraph("Note: All figures are reported in full USD values unless otherwise stated.")
    
    doc.add_heading("Statement of Comprehensive Income", level=2)
    income_data = [
        ("Gross Revenue", "$2,400,000"),
        ("Cost of Goods Sold (COGS)", "$1,440,000"),
        ("Gross Profit", "$960,000"),
        ("Selling, General & Administrative Expenses", "$480,000"),
        ("Operating Income (EBIT)", "$480,000"),
        ("Interest Expense", "$40,000"),
        ("Earnings Before Taxes (EBT)", "$440,000"),
        ("Income Tax Expense", "$88,000"),
        ("Net Income (PAT)", "$352,000"),
    ]
    for item, val in income_data:
        doc.add_paragraph(f"{item}: {val}")
        
    doc.add_heading("Statement of Financial Position (Balance Sheet)", level=2)
    bs_data = [
        ("Cash and Cash Equivalents", "$180,000"),
        ("Trade Receivables", "$320,000"),
        ("Inventories", "$250,000"),
        ("Total Current Assets", "$750,000"),
        ("Property, Plant and Equipment (Net)", "$1,250,000"),
        ("Total Assets", "$2,000,000"),
        ("Trade Payables", "$200,000"),
        ("Short-Term Debt", "$100,000"),
        ("Total Current Liabilities", "$300,000"),
        ("Long-Term Debt", "$500,000"),
        ("Total Liabilities", "$800,000"),
        ("Share Capital", "$600,000"),
        ("Retained Earnings", "$600,000"),
        ("Total Shareholders' Equity", "$1,200,000"),
        ("Total Liabilities and Equity", "$2,000,000"),
    ]
    for item, val in bs_data:
        doc.add_paragraph(f"{item}: {val}")
        
    doc.add_heading("Statement of Cash Flows", level=2)
    cf_data = [
        ("Cash Flow from Operating Activities", "$420,000"),
        ("Cash Flow from Investing Activities", "-$260,000"),
        ("Cash Flow from Financing Activities", "-$80,000"),
        ("Net Increase in Cash and Cash Equivalents", "$80,000"),
        ("Cash at Beginning of Year", "$100,000"),
        ("Cash at End of Year", "$180,000"),
    ]
    for item, val in cf_data:
        doc.add_paragraph(f"{item}: {val}")
        
    doc.add_heading("Analytical Review Questions & Auditor Instructions", level=2)
    questions = [
        "1. What is the Net Profit Margin for AeroPulse Systems Corp. for FY 2024–25?",
        "2. Does the Balance Sheet logically balance (Total Assets == Total Liabilities + Equity)?",
        "3. Does the Ending Cash on the Cash Flow Statement reconcile with Cash on the Balance Sheet?",
        "4. What is the Debt-to-Equity ratio of the corporation?",
        "5. What is the Current Ratio and Quick Ratio?",
    ]
    for q in questions:
        doc.add_paragraph(q)
        
    doc.save(filepath)
    print(f"Created unseen document at: {filepath}")

def run_test():
    filepath = "uploaded_files/unseen_level1_aeropulse.docx"
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    create_unseen_docx(filepath)
    
    with open(filepath, "rb") as f:
        file_bytes = f.read()
        
    print("\n--- STAGE 1: DOCUMENT PARSING ---")
    parsed = parse_workbook(file_bytes, "unseen_level1_aeropulse.docx")
    items = parsed.get("normalized_items", [])
    company_name = parsed.get("company_name", "")
    currency = parsed.get("currency", "NOT_DETERMINED")
    print(f"Extracted Company: {company_name}")
    print(f"Detected Currency: {currency}")
    print(f"Extracted Items Count: {len(items)}")
    
    # Check that questions were NOT extracted as financial rows
    question_leaks = [it for it in items if "?" in it.get("canonical_name", "") or "What is" in it.get("canonical_name", "") or "?" in it.get("original_name", "")]
    print(f"Question leaks count: {len(question_leaks)}")
    assert len(question_leaks) == 0, f"Found question leaks: {question_leaks}"
    
    print("\n--- STAGE 2: FINANCIAL STATEMENTS GENERATION ---")
    statements = sanitize_json_data(generate_financial_statements(items))
    periods = statements.get("periods", [])
    print(f"Statement Periods: {periods}")
    p_key = periods[0] if periods else "FY2025"
    print(f"Active Period Key: {p_key}")
    
    is_stmt = statements.get("income_statement", {})
    bs_stmt = statements.get("balance_sheet", {})
    cf_stmt = statements.get("cash_flow", {})
    
    rev = is_stmt.get("total_revenue") or is_stmt.get("revenue_from_operations") or is_stmt.get("sales", 0.0)
    cogs = is_stmt.get("cost_of_goods_sold") or is_stmt.get("cogs", 0.0)
    gp = is_stmt.get("gross_profit", 0.0)
    ebit = is_stmt.get("ebit") or is_stmt.get("profit_from_operations", 0.0)
    pbt = is_stmt.get("pbt") or is_stmt.get("ebt", 0.0)
    ni = is_stmt.get("net_income") or is_stmt.get("net_profit", 0.0)
    
    print(f"Revenue: {rev}")
    print(f"COGS: {cogs}")
    print(f"Gross Profit: {gp}")
    print(f"Operating Income (EBIT): {ebit}")
    print(f"PBT: {pbt}")
    print(f"Net Income: {ni}")
    
    tot_assets = bs_stmt.get("total_assets", 0.0)
    curr_assets = bs_stmt.get("current_assets", {}).get("total_current_assets", 0.0) if isinstance(bs_stmt.get("current_assets"), dict) else 0.0
    tot_liab = bs_stmt.get("total_liabilities", 0.0)
    tot_equity = bs_stmt.get("equity", {}).get("total_equity", 0.0) if isinstance(bs_stmt.get("equity"), dict) else 0.0
    tot_liab_eq = bs_stmt.get("total_liabilities_and_equity", 0.0)
    
    print(f"Current Assets: {curr_assets}")
    print(f"Total Assets: {tot_assets}")
    print(f"Total Liabilities: {tot_liab}")
    print(f"Total Equity: {tot_equity}")
    print(f"Total Liab & Equity: {tot_liab_eq}")
    
    ocf = cf_stmt.get("operating_activities", 0.0)
    icf = cf_stmt.get("investing_activities", 0.0)
    fcf = cf_stmt.get("financing_activities", 0.0)
    net_cf = cf_stmt.get("net_change_in_cash", 0.0)
    
    print(f"Operating Cash Flow: {ocf}")
    print(f"Investing Cash Flow: {icf}")
    print(f"Financing Cash Flow: {fcf}")
    print(f"Net Cash Flow: {net_cf}")
    
    print("\n--- STAGE 3: RATIOS & CORPORATE FINANCE ---")
    ratios = sanitize_json_data(calculate_financial_ratios(statements))
    corp_fin = sanitize_json_data(calculate_corporate_finance(statements, ratios))
    
    curr_ratio = ratios.get("liquidity", {}).get("current_ratio", {}).get("value")
    quick_ratio = ratios.get("liquidity", {}).get("quick_ratio", {}).get("value")
    npm = ratios.get("profitability", {}).get("net_profit_margin", {}).get("value")
    de_ratio = ratios.get("solvency", {}).get("debt_to_equity", {}).get("value")
    
    print(f"Current Ratio: {curr_ratio}")
    print(f"Quick Ratio: {quick_ratio}")
    print(f"Net Profit Margin: {npm}%")
    print(f"Debt to Equity: {de_ratio}")
    
    print("\n--- STAGE 4: RECONCILIATION & QUALITY ---")
    canonical_dataset = sanitize_json_data(build_canonical_dataset(items, "unseen_level1_aeropulse.docx"))
    reconciliation = sanitize_json_data(perform_source_to_result_reconciliation(canonical_dataset, statements, ratios))
    quality_report = sanitize_json_data(compute_financial_quality_score(reconciliation, statements.get("validation_report", {})))
    
    recon_summary = reconciliation.get("summary", {})
    print(f"Reconciliation Summary: {recon_summary}")
    print(f"Quality Score: {quality_report.get('financial_quality_score')}")
    
    print("\n--- STAGE 5: AI INSIGHTS ---")
    ai_insights = sanitize_json_data(generate_ai_insights(statements, ratios, corp_fin, canonical_dataset, quality_report=quality_report))
    print("Executive Summary:")
    print(ai_insights.get("executive_summary"))
    
    print("\n--- STAGE 6: DETERMINISTIC AUDIT ASSERTIONS ---")
    # 1. Income Statement
    assert rev == 2400000.0, f"Expected Rev 2,400,000, got {rev}"
    assert cogs == 1440000.0, f"Expected COGS 1,440,000, got {cogs}"
    assert gp == 960000.0, f"Expected GP 960,000, got {gp}"
    assert ebit == 480000.0, f"Expected EBIT 480,000, got {ebit}"
    assert pbt == 440000.0, f"Expected PBT 440,000, got {pbt}"
    assert ni == 352000.0, f"Expected Net Income 352,000, got {ni}"
    print("[PASS] Income Statement Math verified: Rev(2.4M) - COGS(1.44M) = GP(960k), EBIT(480k) - Int(40k) = PBT(440k), PBT - Tax(88k) = NI(352k).")
    
    # 2. Balance Sheet
    assert tot_assets == 2000000.0, f"Expected Assets 2,000,000, got {tot_assets}"
    assert tot_liab == 800000.0, f"Expected Liab 800,000, got {tot_liab}"
    assert tot_equity == 1200000.0, f"Expected Equity 1,200,000, got {tot_equity}"
    assert tot_liab_eq == 2000000.0, f"Expected Liab+Eq 2,000,000, got {tot_liab_eq}"
    assert tot_assets == tot_liab_eq, "Balance Sheet does not balance!"
    print("[PASS] Balance Sheet Equilibrium: Assets(2.0M) == Liabilities(800k) + Equity(1.2M).")
    
    # 3. Cash Flow
    assert ocf == 420000.0, f"Expected OCF 420,000, got {ocf}"
    assert icf == -260000.0, f"Expected ICF -260,000, got {icf}"
    assert fcf == -80000.0, f"Expected FCF -80,000, got {fcf}"
    assert net_cf == 80000.0, f"Expected Net CF 80,000, got {net_cf}"
    print("[PASS] Cash Flow Statement: OCF(420k) + ICF(-260k) + FCF(-80k) = Net Increase in Cash(80k).")
    
    # 4. Cash Flow to Balance Sheet link
    cash_bs = bs_stmt.get("current_assets", {}).get("cash", 0.0) if isinstance(bs_stmt.get("current_assets"), dict) else 0.0
    print(f"Balance Sheet Cash: {cash_bs}")
    assert cash_bs == 180000.0, f"Expected BS Cash 180,000, got {cash_bs}"
    print("[PASS] Cash link: Ending Cash on CF statement ($180,000) exactly reconciles with Balance Sheet Cash ($180,000).")
    
    # 5. Ratios
    # NPM = 352,000 / 2,400,000 = 14.67%
    assert round(npm, 2) == 14.67, f"Expected NPM 14.67%, got {npm}"
    # Current Ratio = 750,000 / 300,000 = 2.50
    assert round(curr_ratio, 2) == 2.50, f"Expected CR 2.50, got {curr_ratio}"
    # Quick Ratio = (750,000 - 250,000) / 300,000 = 500,000 / 300,000 = 1.67
    assert round(quick_ratio, 2) == 1.67, f"Expected QR 1.67, got {quick_ratio}"
    # Debt-to-Equity (Interest Bearing = 100k + 500k = 600k; 600k / 1.2M = 0.50)
    assert round(de_ratio, 2) == 0.50, f"Expected D/E 0.50, got {de_ratio}"
    print("[PASS] Ratios verified: NPM=14.67%, CR=2.50, QR=1.67, D/E=0.50.")
    
    print("\nSUCCESS: Unseen test case passed all financial reconciliations with ZERO errors!")

if __name__ == "__main__":
    run_test()
